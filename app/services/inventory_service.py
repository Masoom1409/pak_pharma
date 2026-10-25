"""Medicine Master, Barcode Scanning, FEFO Batch Management, and CSV/Excel Import-Export Service."""

from __future__ import annotations

import csv
from datetime import date, datetime, timedelta
from decimal import Decimal
import os
import re
import sqlite3
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from app.db.database import DatabaseManager, format_pkr, paisa_to_pkr, pkr_to_paisa
from app.services.auth_service import AuthService


class BarcodeValidationError(ValueError):
    """Raised when a barcode string is invalid, empty, or already assigned."""


class DuplicateProductError(ValueError):
    """Raised when attempting to create an accidental duplicate medicine."""


class InsufficientStockError(ValueError):
    """Raised when requested sale or adjustment quantity exceeds available stock."""


class ExpiredStockError(ValueError):
    """Raised when attempting to sell expired stock without authorized override."""


def compute_gs1_check_digit(digits_without_check: str) -> str:
    """Compute standard GS1 modulo-10 check digit for EAN-13, EAN-8, or UPC-A."""
    total = 0
    # From rightmost digit moving left: weights alternate 3, 1, 3, 1...
    for idx, ch in enumerate(reversed(digits_without_check)):
        weight = 3 if idx % 2 == 0 else 1
        total += int(ch) * weight
    remainder = total % 10
    return str((10 - remainder) % 10)


def validate_and_classify_barcode(barcode: str) -> tuple[str, str]:
    """Validate barcode string and classify as EAN13, EAN8, UPCA, or CODE128."""
    clean = (barcode or "").strip()
    if not clean:
        raise BarcodeValidationError("Barcode cannot be empty.")
    if len(clean) < 4 or len(clean) > 48:
        raise BarcodeValidationError("Barcode length must be between 4 and 48 characters.")
    if not re.match(r"^[A-Za-z0-9\-_./]+$", clean):
        raise BarcodeValidationError(
            f"Invalid characters in barcode '{clean}'. Only alphanumeric and [-_./] are allowed."
        )

    if clean.isdigit():
        if len(clean) == 13:
            expected = compute_gs1_check_digit(clean[:12])
            if clean[12] != expected:
                raise BarcodeValidationError(
                    f"Invalid EAN-13 check digit in '{clean}' (expected '{expected}', got '{clean[12]}')."
                )
            return clean, "EAN13"
        if len(clean) == 8:
            expected = compute_gs1_check_digit(clean[:7])
            if clean[7] != expected:
                raise BarcodeValidationError(
                    f"Invalid EAN-8 check digit in '{clean}' (expected '{expected}', got '{clean[7]}')."
                )
            return clean, "EAN8"
        if len(clean) == 12:
            expected = compute_gs1_check_digit(clean[:11])
            if clean[11] != expected:
                raise BarcodeValidationError(
                    f"Invalid UPC-A check digit in '{clean}' (expected '{expected}', got '{clean[11]}')."
                )
            return clean, "UPCA"

    return clean, "CODE128"


def generate_internal_ean13(product_id: int) -> str:
    """Generate a valid in-store EAN-13 barcode (prefix 200 for internal store use)."""
    base_12 = f"200{product_id:09d}"[:12]
    check = compute_gs1_check_digit(base_12)
    return base_12 + check


class InventoryService:
    """Manages Medicine Master, Barcodes, Batches, FEFO allocation, and Stock Adjustments."""

    def __init__(self, db: DatabaseManager, auth: AuthService) -> None:
        self.db = db
        self.auth = auth

    def _next_product_code(self, conn: sqlite3.Connection) -> str:
        row = conn.execute("SELECT COALESCE(MAX(id), 0) + 1 AS next_id FROM products").fetchone()
        return f"MED-{int(row['next_id']):05d}"

    def create_product(
        self,
        user_id: int,
        medicine_name: str,
        generic_name: str,
        manufacturer: str,
        purchase_price_pkr: Decimal | int | float | str,
        sale_price_pkr: Decimal | int | float | str,
        brand_name: str = "",
        category: str = "General",
        dosage_form: str = "Tablet",
        strength: str = "",
        pack_size: int = 1,
        unit_of_measure: str = "Tablet",
        pack_unit_name: str = "Box",
        barcode: str | None = None,
        alternate_barcodes: list[tuple[str, int]] | None = None,
        min_stock_threshold: int = 10,
        preferred_supplier_id: int | None = None,
        initial_batch_number: str | None = None,
        initial_expiry_date: str | None = None,
        initial_qty: int = 0,
        is_sample_record: bool = False,
        verification_status: str = "USER_ENTERED",
        notes: str = "",
        product_code: str | None = None,
    ) -> dict[str, Any]:
        """Create a medicine master record, assign primary/alternate barcodes, and optionally add initial batch."""
        actor = self.auth.get_user(user_id)
        med_name = (medicine_name or "").strip()
        gen_name = (generic_name or "").strip()
        mfg_name = (manufacturer or "").strip()
        strength_clean = (strength or "").strip()
        dosage_clean = (dosage_form or "Tablet").strip()

        if not med_name:
            raise ValueError("Medicine name is required.")
        if not gen_name:
            raise ValueError("Generic / salt name is required.")
        if not mfg_name:
            raise ValueError("Manufacturer name is required.")
        if pack_size < 1:
            raise ValueError("Pack size must be at least 1.")

        p_paisa = pkr_to_paisa(purchase_price_pkr)
        s_paisa = pkr_to_paisa(sale_price_pkr)
        if p_paisa < 0 or s_paisa < 0:
            raise ValueError("Prices cannot be negative.")

        validated_primary: tuple[str, str] | None = None
        if barcode and barcode.strip():
            validated_primary = validate_and_classify_barcode(barcode)

        validated_alternates: list[tuple[str, str, int]] = []
        if alternate_barcodes:
            for alt_code, mult in alternate_barcodes:
                if alt_code and alt_code.strip():
                    c_code, c_type = validate_and_classify_barcode(alt_code)
                    validated_alternates.append((c_code, c_type, max(1, int(mult))))

        now_iso = datetime.now().isoformat(timespec="seconds")

        with self.db.transaction() as conn:
            # Check for duplicate medicine (name + strength + dosage_form + manufacturer + pack_size)
            dup = conn.execute(
                """
                SELECT id, product_code FROM products
                WHERE LOWER(medicine_name) = LOWER(?)
                  AND LOWER(strength) = LOWER(?)
                  AND LOWER(dosage_form) = LOWER(?)
                  AND LOWER(manufacturer) = LOWER(?)
                  AND pack_size = ?
                """,
                (med_name, strength_clean, dosage_clean, mfg_name, int(pack_size)),
            ).fetchone()
            if dup:
                raise DuplicateProductError(
                    f"Duplicate medicine exists: '{med_name} {strength_clean} ({dosage_clean})' "
                    f"by {mfg_name} [Code: {dup['product_code']}]."
                )

            # Check barcode uniqueness before inserting
            all_codes_to_check = []
            if validated_primary:
                all_codes_to_check.append(validated_primary[0])
            for ac, _, _ in validated_alternates:
                if ac in all_codes_to_check:
                    raise BarcodeValidationError(f"Duplicate barcode '{ac}' provided in the same product.")
                all_codes_to_check.append(ac)

            for code_str in all_codes_to_check:
                existing_bc = conn.execute(
                    """
                    SELECT pb.barcode, p.medicine_name, p.product_code
                    FROM product_barcodes pb
                    JOIN products p ON p.id = pb.product_id
                    WHERE LOWER(pb.barcode) = LOWER(?)
                    """,
                    (code_str,),
                ).fetchone()
                if existing_bc:
                    raise BarcodeValidationError(
                        f"Barcode '{code_str}' is already assigned to '{existing_bc['medicine_name']}' "
                        f"({existing_bc['product_code']})."
                    )

            code = (product_code or "").strip() or self._next_product_code(conn)
            cur = conn.execute(
                """
                INSERT INTO products (
                    product_code, medicine_name, generic_name, brand_name, manufacturer,
                    category, dosage_form, strength, pack_size, unit_of_measure, pack_unit_name,
                    default_purchase_price_paisa, default_sale_price_paisa, min_stock_threshold,
                    preferred_supplier_id, is_active, is_sample_record, verification_status,
                    notes, created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1, ?, ?, ?, ?, ?)
                """,
                (
                    code,
                    med_name,
                    gen_name,
                    (brand_name or med_name).strip(),
                    mfg_name,
                    (category or "General").strip(),
                    dosage_clean,
                    strength_clean,
                    int(pack_size),
                    (unit_of_measure or "Unit").strip(),
                    (pack_unit_name or "Pack").strip(),
                    p_paisa,
                    s_paisa,
                    int(min_stock_threshold),
                    preferred_supplier_id,
                    1 if is_sample_record else 0,
                    verification_status,
                    notes.strip(),
                    now_iso,
                    now_iso,
                ),
            )
            product_id = int(cur.lastrowid)

            # If no barcode was supplied, generate an internal EAN-13 barcode
            if validated_primary:
                primary_code, primary_type = validated_primary
            else:
                primary_code = generate_internal_ean13(product_id)
                primary_type = "EAN13"

            conn.execute(
                """
                INSERT INTO product_barcodes (product_id, barcode, barcode_type, is_primary, pack_multiplier, created_at)
                VALUES (?, ?, ?, 1, 1, ?)
                """,
                (product_id, primary_code, primary_type, now_iso),
            )

            for alt_code, alt_type, mult in validated_alternates:
                conn.execute(
                    """
                    INSERT INTO product_barcodes (product_id, barcode, barcode_type, is_primary, pack_multiplier, created_at)
                    VALUES (?, ?, ?, 0, ?, ?)
                    """,
                    (product_id, alt_code, alt_type, mult, now_iso),
                )

            batch_id = None
            if initial_batch_number and initial_expiry_date and int(initial_qty) >= 0:
                self._validate_date_str(initial_expiry_date)
                b_cur = conn.execute(
                    """
                    INSERT INTO batches (
                        product_id, batch_number, expiry_date, qty_received, qty_remaining,
                        purchase_cost_paisa, sale_price_paisa, supplier_id, purchase_reference,
                        is_quarantined, created_at, updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'INITIAL_OPENING', 0, ?, ?)
                    """,
                    (
                        product_id,
                        initial_batch_number.strip(),
                        initial_expiry_date.strip(),
                        int(initial_qty),
                        int(initial_qty),
                        p_paisa,
                        s_paisa,
                        preferred_supplier_id,
                        now_iso,
                        now_iso,
                    ),
                )
                batch_id = int(b_cur.lastrowid)
                if int(initial_qty) > 0:
                    conn.execute(
                        """
                        INSERT INTO stock_movements (
                            timestamp, product_id, batch_id, movement_type, qty_change,
                            qty_before, qty_after, unit_cost_paisa, reference_type, reference_id,
                            user_id, notes
                        ) VALUES (?, ?, ?, 'INITIAL_STOCK', ?, 0, ?, ?, 'PRODUCT_INIT', ?, ?, ?)
                        """,
                        (
                            now_iso,
                            product_id,
                            batch_id,
                            int(initial_qty),
                            int(initial_qty),
                            p_paisa,
                            code,
                            user_id,
                            f"Initial stock for batch {initial_batch_number}",
                        ),
                    )

            self.auth._log_in_tx(
                conn,
                user_id=actor["id"],
                username=actor["username"],
                action="CREATE_PRODUCT",
                entity_type="products",
                entity_id=str(product_id),
                reason=f"Created medicine {med_name} ({code})",
                details={"barcode": primary_code, "initial_qty": initial_qty},
            )

        return self.get_product_by_id(product_id)

    def update_product(
        self,
        user_id: int,
        product_id: int,
        medicine_name: str | None = None,
        generic_name: str | None = None,
        brand_name: str | None = None,
        manufacturer: str | None = None,
        category: str | None = None,
        dosage_form: str | None = None,
        strength: str | None = None,
        pack_size: int | None = None,
        unit_of_measure: str | None = None,
        pack_unit_name: str | None = None,
        purchase_price_pkr: Decimal | int | float | str | None = None,
        sale_price_pkr: Decimal | int | float | str | None = None,
        min_stock_threshold: int | None = None,
        preferred_supplier_id: int | None = None,
        is_active: bool | None = None,
        verification_status: str | None = None,
        notes: str | None = None,
        reason: str = "Updated medicine details",
    ) -> dict[str, Any]:
        """Update medicine details. Enforces CHANGE_PRICES permission if purchase or sale price changes."""
        actor = self.auth.get_user(user_id)
        existing = self.get_product_by_id(product_id)

        new_p_paisa = (
            pkr_to_paisa(purchase_price_pkr)
            if purchase_price_pkr is not None
            else existing["default_purchase_price_paisa"]
        )
        new_s_paisa = (
            pkr_to_paisa(sale_price_pkr)
            if sale_price_pkr is not None
            else existing["default_sale_price_paisa"]
        )

        if (
            new_p_paisa != existing["default_purchase_price_paisa"]
            or new_s_paisa != existing["default_sale_price_paisa"]
        ):
            self.auth.require_permission(user_id, "CHANGE_PRICES", reason="Attempted to change medicine price")

        if is_active is not None and bool(is_active) != bool(existing["is_active"]):
            self.auth.require_permission(user_id, "DEACTIVATE_PRODUCT", reason="Attempted to change product status")

        now_iso = datetime.now().isoformat(timespec="seconds")
        with self.db.transaction() as conn:
            conn.execute(
                """
                UPDATE products SET
                    medicine_name = ?,
                    generic_name = ?,
                    brand_name = ?,
                    manufacturer = ?,
                    category = ?,
                    dosage_form = ?,
                    strength = ?,
                    pack_size = ?,
                    unit_of_measure = ?,
                    pack_unit_name = ?,
                    default_purchase_price_paisa = ?,
                    default_sale_price_paisa = ?,
                    min_stock_threshold = ?,
                    preferred_supplier_id = ?,
                    is_active = ?,
                    verification_status = ?,
                    notes = ?,
                    updated_at = ?
                WHERE id = ?
                """,
                (
                    (medicine_name if medicine_name is not None else existing["medicine_name"]).strip(),
                    (generic_name if generic_name is not None else existing["generic_name"]).strip(),
                    (brand_name if brand_name is not None else existing["brand_name"]).strip(),
                    (manufacturer if manufacturer is not None else existing["manufacturer"]).strip(),
                    (category if category is not None else existing["category"]).strip(),
                    (dosage_form if dosage_form is not None else existing["dosage_form"]).strip(),
                    (strength if strength is not None else existing["strength"]).strip(),
                    int(pack_size if pack_size is not None else existing["pack_size"]),
                    (unit_of_measure if unit_of_measure is not None else existing["unit_of_measure"]).strip(),
                    (pack_unit_name if pack_unit_name is not None else existing["pack_unit_name"]).strip(),
                    new_p_paisa,
                    new_s_paisa,
                    int(min_stock_threshold if min_stock_threshold is not None else existing["min_stock_threshold"]),
                    preferred_supplier_id if preferred_supplier_id is not None else existing["preferred_supplier_id"],
                    1 if (is_active if is_active is not None else existing["is_active"]) else 0,
                    verification_status if verification_status is not None else existing["verification_status"],
                    (notes if notes is not None else (existing["notes"] or "")).strip(),
                    now_iso,
                    product_id,
                ),
            )
            self.auth._log_in_tx(
                conn,
                user_id=actor["id"],
                username=actor["username"],
                action="UPDATE_PRODUCT",
                entity_type="products",
                entity_id=str(product_id),
                reason=reason,
                details={
                    "old_sale_paisa": existing["default_sale_price_paisa"],
                    "new_sale_paisa": new_s_paisa,
                },
            )
        return self.get_product_by_id(product_id)

    def add_barcode_to_product(
        self,
        user_id: int,
        product_id: int,
        barcode: str,
        is_primary: bool = False,
        pack_multiplier: int = 1,
    ) -> dict[str, Any]:
        """Attach an additional or replacement barcode to an existing product."""
        actor = self.auth.get_user(user_id)
        clean_code, code_type = validate_and_classify_barcode(barcode)
        now_iso = datetime.now().isoformat(timespec="seconds")
        with self.db.transaction() as conn:
            dup = conn.execute(
                "SELECT product_id FROM product_barcodes WHERE LOWER(barcode) = LOWER(?)",
                (clean_code,),
            ).fetchone()
            if dup:
                raise BarcodeValidationError(f"Barcode '{clean_code}' is already registered.")
            if is_primary:
                conn.execute(
                    "UPDATE product_barcodes SET is_primary = 0 WHERE product_id = ?",
                    (product_id,),
                )
            conn.execute(
                """
                INSERT INTO product_barcodes (product_id, barcode, barcode_type, is_primary, pack_multiplier, created_at)
                VALUES (?, ?, ?, ?, ?, ?)
                """,
                (product_id, clean_code, code_type, 1 if is_primary else 0, max(1, int(pack_multiplier)), now_iso),
            )
            self.auth._log_in_tx(
                conn,
                user_id=actor["id"],
                username=actor["username"],
                action="ADD_BARCODE",
                entity_type="product_barcodes",
                entity_id=clean_code,
                reason=f"Added barcode {clean_code} to product {product_id}",
            )
        return self.get_product_by_id(product_id)

    def get_product_by_id(self, product_id: int) -> dict[str, Any]:
        """Retrieve complete medicine record including barcodes, batches, and total stock."""
        with self.db.read_only() as conn:
            row = conn.execute(
                """
                SELECT p.*,
                       COALESCE(SUM(CASE WHEN b.is_quarantined = 0 THEN b.qty_remaining ELSE 0 END), 0) AS total_stock,
                       s.name AS supplier_name
                FROM products p
                LEFT JOIN batches b ON b.product_id = p.id
                LEFT JOIN suppliers s ON s.id = p.preferred_supplier_id
                WHERE p.id = ?
                GROUP BY p.id
                """,
                (product_id,),
            ).fetchone()
            if not row:
                raise ValueError(f"Product ID {product_id} not found.")
            prod = dict(row)
            bcs = conn.execute(
                "SELECT barcode, barcode_type, is_primary, pack_multiplier FROM product_barcodes WHERE product_id = ? ORDER BY is_primary DESC, id ASC",
                (product_id,),
            ).fetchall()
            prod["barcodes"] = [dict(b) for b in bcs]
            prod["primary_barcode"] = bcs[0]["barcode"] if bcs else ""
            batches = conn.execute(
                """
                SELECT * FROM batches
                WHERE product_id = ?
                ORDER BY expiry_date ASC, id ASC
                """,
                (product_id,),
            ).fetchall()
            prod["batches"] = [dict(b) for b in batches]
            return prod

    def scan_barcode(self, scanned_code: str, reference_date: str | None = None) -> dict[str, Any]:
        """Fast barcode lookup for USB keyboard-wedge barcode scanners.

        IMPORTANT SAFEGUARD: Matches ONLY exact registered barcode in `product_barcodes`.
        Never auto-selects a product based on name similarity.
        If found, returns the product and automatically selects the best FEFO batch.
        If unknown, returns `found=False` with `prefill_barcode` so UI opens New Medicine dialog.
        """
        clean = (scanned_code or "").strip()
        if not clean:
            raise BarcodeValidationError("Empty barcode scan received.")

        ref_date = reference_date or date.today().isoformat()
        with self.db.read_only() as conn:
            bc_row = conn.execute(
                """
                SELECT pb.product_id, pb.barcode, pb.barcode_type, pb.is_primary, pb.pack_multiplier
                FROM product_barcodes pb
                JOIN products p ON p.id = pb.product_id
                WHERE LOWER(pb.barcode) = LOWER(?) AND p.is_active = 1
                """,
                (clean,),
            ).fetchone()

            if not bc_row:
                return {
                    "found": False,
                    "prefill_barcode": clean,
                    "product": None,
                    "selected_batch": None,
                    "available_batches": [],
                    "expired_batches": [],
                    "pack_multiplier": 1,
                }

            product = self.get_product_by_id(int(bc_row["product_id"]))
            valid_batches: list[dict[str, Any]] = []
            expired_batches: list[dict[str, Any]] = []
            for b in product["batches"]:
                if b["qty_remaining"] > 0 and not b["is_quarantined"]:
                    if b["expiry_date"] >= ref_date:
                        valid_batches.append(b)
                    else:
                        expired_batches.append(b)

            # FEFO: valid_batches is already ordered by expiry_date ASC, id ASC
            selected_batch = valid_batches[0] if valid_batches else None

            return {
                "found": True,
                "prefill_barcode": clean,
                "scanned_barcode": bc_row["barcode"],
                "barcode_type": bc_row["barcode_type"],
                "pack_multiplier": int(bc_row["pack_multiplier"]),
                "product": product,
                "selected_batch": selected_batch,
                "available_batches": valid_batches,
                "expired_batches": expired_batches,
            }

    def search_products(
        self,
        query: str = "",
        category: str = "",
        manufacturer: str = "",
        stock_filter: str = "ALL",  # ALL, LOW_STOCK, OUT_OF_STOCK, IN_STOCK
        expiry_filter: str = "ALL",  # ALL, EXPIRED, NEAR_EXPIRY_30, NEAR_EXPIRY_60, NEAR_EXPIRY_90
        include_inactive: bool = False,
        reference_date: str | None = None,
    ) -> list[dict[str, Any]]:
        """Search and filter medicines by name, generic, brand, barcode, code, category, stock, and expiry."""
        ref_date = reference_date or date.today().isoformat()
        d_ref = date.fromisoformat(ref_date)

        with self.db.read_only() as conn:
            sql = """
                SELECT p.*,
                       COALESCE((
                           SELECT pb.barcode FROM product_barcodes pb
                           WHERE pb.product_id = p.id
                           ORDER BY pb.is_primary DESC, pb.id ASC LIMIT 1
                       ), '') AS primary_barcode,
                       COALESCE(SUM(CASE WHEN b.is_quarantined = 0 AND b.expiry_date >= ? THEN b.qty_remaining ELSE 0 END), 0) AS sellable_stock,
                       COALESCE(SUM(CASE WHEN b.is_quarantined = 0 THEN b.qty_remaining ELSE 0 END), 0) AS total_stock,
                       COALESCE(SUM(CASE WHEN b.qty_remaining > 0 AND b.expiry_date < ? THEN b.qty_remaining ELSE 0 END), 0) AS expired_stock,
                       MIN(CASE WHEN b.qty_remaining > 0 AND b.is_quarantined = 0 AND b.expiry_date >= ? THEN b.expiry_date ELSE NULL END) AS earliest_valid_expiry,
                       MIN(CASE WHEN b.qty_remaining > 0 THEN b.expiry_date ELSE NULL END) AS earliest_any_expiry
                FROM products p
                LEFT JOIN batches b ON b.product_id = p.id
                WHERE 1=1
            """
            params: list[Any] = [ref_date, ref_date, ref_date]

            if not include_inactive:
                sql += " AND p.is_active = 1"

            if query and query.strip():
                q = f"%{query.strip()}%"
                sql += """
                    AND (
                        p.medicine_name LIKE ?
                        OR p.generic_name LIKE ?
                        OR p.brand_name LIKE ?
                        OR p.product_code LIKE ?
                        OR EXISTS (
                            SELECT 1 FROM product_barcodes pb2
                            WHERE pb2.product_id = p.id AND pb2.barcode LIKE ?
                        )
                    )
                """
                params.extend([q, q, q, q, q])

            if category and category.strip() and category != "ALL":
                sql += " AND LOWER(p.category) = LOWER(?)"
                params.append(category.strip())

            if manufacturer and manufacturer.strip() and manufacturer != "ALL":
                sql += " AND LOWER(p.manufacturer) = LOWER(?)"
                params.append(manufacturer.strip())

            sql += " GROUP BY p.id ORDER BY p.medicine_name ASC"
            rows = [dict(r) for r in conn.execute(sql, params).fetchall()]

        # Apply post-aggregation stock and expiry filters
        filtered: list[dict[str, Any]] = []
        for r in rows:
            if stock_filter == "LOW_STOCK" and r["sellable_stock"] > r["min_stock_threshold"]:
                continue
            if stock_filter == "OUT_OF_STOCK" and r["sellable_stock"] > 0:
                continue
            if stock_filter == "IN_STOCK" and r["sellable_stock"] <= 0:
                continue

            if expiry_filter == "EXPIRED" and r["expired_stock"] <= 0:
                continue
            if expiry_filter.startswith("NEAR_EXPIRY_"):
                days = int(expiry_filter.split("_")[-1])
                horizon = (d_ref + timedelta(days=days)).isoformat()
                earliest = r["earliest_valid_expiry"]
                if not earliest or earliest > horizon:
                    continue

            filtered.append(r)

        return filtered

    def allocate_fefo_batches(
        self,
        conn: sqlite3.Connection,
        product_id: int,
        required_qty: int,
        reference_date: str | None = None,
        allow_expired_override: bool = False,
        override_user_id: int | None = None,
        override_reason: str = "",
        specific_batch_id: int | None = None,
    ) -> list[dict[str, Any]]:
        """Select stock batches using strict First-Expiry-First-Out (FEFO) rules.

        - Excludes quarantined batches.
        - Excludes expired batches unless `allow_expired_override=True` AND user has `EXPIRED_OVERRIDE` permission
          AND a non-empty `override_reason` is provided.
        - Splits across multiple batches in ascending expiry order if a single batch has less than `required_qty`.
        - Raises `ExpiredStockError` if only expired stock exists and override is not active.
        - Raises `InsufficientStockError` if total eligible stock is less than `required_qty`.
        """
        if required_qty <= 0:
            raise ValueError("Required quantity must be greater than zero.")

        ref_date = reference_date or date.today().isoformat()
        block_expired = self.db.get_setting("block_expired_sales", "1") == "1"

        if allow_expired_override:
            if self.db.get_setting("allow_admin_expired_override", "1") != "1":
                raise ExpiredStockError("Expired stock override is disabled in store settings.")
            if not override_user_id:
                raise ExpiredStockError("Authorized user ID is required for expired stock override.")
            if not (override_reason or "").strip():
                raise ExpiredStockError("A written reason is required to override expired stock restriction.")
            self.auth.require_permission(
                override_user_id,
                "EXPIRED_OVERRIDE",
                reason=f"Expired override for product {product_id}: {override_reason}",
            )

        if specific_batch_id is not None:
            batches = conn.execute(
                """
                SELECT * FROM batches
                WHERE id = ? AND product_id = ? AND qty_remaining > 0
                """,
                (specific_batch_id, product_id),
            ).fetchall()
        else:
            batches = conn.execute(
                """
                SELECT * FROM batches
                WHERE product_id = ? AND qty_remaining > 0 AND is_quarantined = 0
                ORDER BY expiry_date ASC, id ASC
                """,
                (product_id,),
            ).fetchall()

        eligible_batches: list[sqlite3.Row] = []
        expired_available_qty = 0

        for b in batches:
            if b["is_quarantined"]:
                continue
            is_expired = b["expiry_date"] < ref_date
            if is_expired and block_expired and not allow_expired_override:
                expired_available_qty += int(b["qty_remaining"])
                continue
            eligible_batches.append(b)

        total_eligible = sum(int(b["qty_remaining"]) for b in eligible_batches)
        if total_eligible < required_qty:
            prod_row = conn.execute(
                "SELECT medicine_name, product_code FROM products WHERE id = ?",
                (product_id,),
            ).fetchone()
            p_label = f"{prod_row['medicine_name']} ({prod_row['product_code']})" if prod_row else f"ID {product_id}"
            if expired_available_qty > 0 and total_eligible + expired_available_qty >= required_qty:
                raise ExpiredStockError(
                    f"Cannot sell '{p_label}': {expired_available_qty} units in stock are EXPIRED "
                    f"(sellable non-expired stock: {total_eligible}, requested: {required_qty})."
                )
            raise InsufficientStockError(
                f"Insufficient stock for '{p_label}': requested {required_qty} units, "
                f"but only {total_eligible} sellable units available."
            )

        allocations: list[dict[str, Any]] = []
        remaining_to_allocate = required_qty
        for b in eligible_batches:
            if remaining_to_allocate <= 0:
                break
            take_qty = min(remaining_to_allocate, int(b["qty_remaining"]))
            is_expired = b["expiry_date"] < ref_date
            allocations.append(
                {
                    "batch_id": int(b["id"]),
                    "batch_number": b["batch_number"],
                    "expiry_date": b["expiry_date"],
                    "qty_allocated": take_qty,
                    "purchase_cost_paisa": int(b["purchase_cost_paisa"]),
                    "sale_price_paisa": int(b["sale_price_paisa"]),
                    "is_expired": is_expired,
                    "expired_override_by": override_user_id if is_expired else None,
                    "expired_override_reason": override_reason if is_expired else None,
                }
            )
            remaining_to_allocate -= take_qty

        return allocations

    def adjust_batch_stock(
        self,
        user_id: int,
        batch_id: int,
        adjustment_type: str,
        qty_delta: int,
        reason: str,
        notes: str = "",
        adjustment_date: str | None = None,
    ) -> dict[str, Any]:
        """Perform a controlled batch-level stock adjustment or quarantine/write-off with audit trail."""
        actor = self.auth.require_permission(
            user_id,
            "STOCK_ADJUSTMENT",
            reason=f"Stock adjustment ({adjustment_type}) on batch {batch_id}",
        )
        if not (reason or "").strip():
            raise ValueError("A clear reason is required for all stock adjustments and write-offs.")

        valid_types = {
            "DAMAGE_WRITE_OFF",
            "EXPIRY_WRITE_OFF",
            "LOST_STOCK",
            "MANUAL_ADJUSTMENT_IN",
            "MANUAL_ADJUSTMENT_OUT",
            "QUARANTINE_HOLD",
            "QUARANTINE_RELEASE",
        }
        if adjustment_type not in valid_types:
            raise ValueError(f"Invalid adjustment type: {adjustment_type}")

        now_iso = datetime.now().isoformat(timespec="seconds")
        adj_date = adjustment_date or date.today().isoformat()

        with self.db.transaction() as conn:
            batch = conn.execute("SELECT * FROM batches WHERE id = ?", (batch_id,)).fetchone()
            if not batch:
                raise ValueError(f"Batch ID {batch_id} not found.")

            product_id = int(batch["product_id"])
            qty_before = int(batch["qty_remaining"])
            unit_cost = int(batch["purchase_cost_paisa"])

            if adjustment_type == "QUARANTINE_HOLD":
                conn.execute(
                    "UPDATE batches SET is_quarantined = 1, quarantine_reason = ?, updated_at = ? WHERE id = ?",
                    (reason.strip(), now_iso, batch_id),
                )
                effective_delta = 0
                qty_after = qty_before
            elif adjustment_type == "QUARANTINE_RELEASE":
                conn.execute(
                    "UPDATE batches SET is_quarantined = 0, quarantine_reason = NULL, updated_at = ? WHERE id = ?",
                    (now_iso, batch_id),
                )
                effective_delta = 0
                qty_after = qty_before
            else:
                if adjustment_type in ("DAMAGE_WRITE_OFF", "EXPIRY_WRITE_OFF", "LOST_STOCK", "MANUAL_ADJUSTMENT_OUT"):
                    effective_delta = -abs(int(qty_delta))
                else:
                    effective_delta = abs(int(qty_delta))

                if effective_delta == 0:
                    raise ValueError("Adjustment quantity delta must be non-zero.")

                qty_after = qty_before + effective_delta
                if qty_after < 0:
                    raise InsufficientStockError(
                        f"Cannot reduce batch '{batch['batch_number']}' below 0 "
                        f"(current remaining: {qty_before}, adjustment: {effective_delta})."
                    )

                conn.execute(
                    "UPDATE batches SET qty_remaining = ?, updated_at = ? WHERE id = ?",
                    (qty_after, now_iso, batch_id),
                )

            next_num_row = conn.execute("SELECT COALESCE(MAX(id), 0) + 1 AS next_id FROM stock_adjustments").fetchone()
            adj_number = f"ADJ-{int(next_num_row['next_id']):05d}"
            total_cost_impact = abs(effective_delta) * unit_cost

            cur = conn.execute(
                """
                INSERT INTO stock_adjustments (
                    adjustment_number, adjustment_date, timestamp, product_id, batch_id,
                    adjustment_type, qty_delta, unit_cost_paisa, total_cost_impact_paisa,
                    reason, authorized_by, notes
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    adj_number,
                    adj_date,
                    now_iso,
                    product_id,
                    batch_id,
                    adjustment_type,
                    effective_delta,
                    unit_cost,
                    total_cost_impact,
                    reason.strip(),
                    user_id,
                    notes.strip(),
                ),
            )
            adj_id = int(cur.lastrowid)

            conn.execute(
                """
                INSERT INTO stock_movements (
                    timestamp, product_id, batch_id, movement_type, qty_change,
                    qty_before, qty_after, unit_cost_paisa, reference_type, reference_id,
                    user_id, notes
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'STOCK_ADJUSTMENT', ?, ?, ?)
                """,
                (
                    now_iso,
                    product_id,
                    batch_id,
                    adjustment_type,
                    effective_delta,
                    qty_before,
                    qty_after,
                    unit_cost,
                    adj_number,
                    user_id,
                    f"{reason.strip()} | {notes.strip()}".strip(" |"),
                ),
            )

            self.auth._log_in_tx(
                conn,
                user_id=actor["id"],
                username=actor["username"],
                action=f"STOCK_ADJUSTMENT:{adjustment_type}",
                entity_type="batches",
                entity_id=str(batch_id),
                reason=reason.strip(),
                details={
                    "adjustment_number": adj_number,
                    "qty_before": qty_before,
                    "qty_after": qty_after,
                    "cost_impact_paisa": total_cost_impact,
                },
            )

            return {
                "id": adj_id,
                "adjustment_number": adj_number,
                "batch_id": batch_id,
                "product_id": product_id,
                "adjustment_type": adjustment_type,
                "qty_before": qty_before,
                "qty_after": qty_after,
                "qty_delta": effective_delta,
                "total_cost_impact_paisa": total_cost_impact,
            }

    def get_expiry_and_quarantine_report(self, reference_date: str | None = None) -> dict[str, list[dict[str, Any]]]:
        """Return categorized batch lists: expired, within 30 days, 60 days, 90 days, and quarantined."""
        ref_str = reference_date or date.today().isoformat()
        d_ref = date.fromisoformat(ref_str)
        d_30 = (d_ref + timedelta(days=int(self.db.get_setting("expiry_alert_days_1", "30")))).isoformat()
        d_60 = (d_ref + timedelta(days=int(self.db.get_setting("expiry_alert_days_2", "60")))).isoformat()
        d_90 = (d_ref + timedelta(days=int(self.db.get_setting("expiry_alert_days_3", "90")))).isoformat()

        with self.db.read_only() as conn:
            rows = conn.execute(
                """
                SELECT b.*, p.product_code, p.medicine_name, p.generic_name, p.strength,
                       p.dosage_form, p.manufacturer, s.name AS supplier_name
                FROM batches b
                JOIN products p ON p.id = b.product_id
                LEFT JOIN suppliers s ON s.id = b.supplier_id
                WHERE b.qty_remaining > 0
                ORDER BY b.expiry_date ASC, b.id ASC
                """
            ).fetchall()

        expired: list[dict[str, Any]] = []
        near_30: list[dict[str, Any]] = []
        near_60: list[dict[str, Any]] = []
        near_90: list[dict[str, Any]] = []
        quarantined: list[dict[str, Any]] = []

        for r in rows:
            item = dict(r)
            exp_d = date.fromisoformat(item["expiry_date"])
            item["days_to_expiry"] = (exp_d - d_ref).days
            item["stock_value_paisa"] = item["qty_remaining"] * item["purchase_cost_paisa"]

            if item["is_quarantined"]:
                quarantined.append(item)
            if item["expiry_date"] < ref_str:
                expired.append(item)
            elif item["expiry_date"] <= d_30:
                near_30.append(item)
            elif item["expiry_date"] <= d_60:
                near_60.append(item)
            elif item["expiry_date"] <= d_90:
                near_90.append(item)

        return {
            "expired": expired,
            "near_30_days": near_30,
            "near_60_days": near_60,
            "near_90_days": near_90,
            "quarantined": quarantined,
        }

    def export_inventory_to_csv(self, file_path: str) -> str:
        """Export current medicine master and aggregated batch stock to CSV."""
        products = self.search_products(include_inactive=True)
        os.makedirs(os.path.dirname(os.path.abspath(file_path)), exist_ok=True)
        headers = [
            "product_code",
            "medicine_name",
            "generic_name",
            "brand_name",
            "manufacturer",
            "category",
            "dosage_form",
            "strength",
            "pack_size",
            "unit_of_measure",
            "primary_barcode",
            "purchase_price_pkr",
            "sale_price_pkr",
            "min_stock_threshold",
            "sellable_stock",
            "expired_stock",
            "earliest_valid_expiry",
            "verification_status",
            "is_sample_record",
            "is_active",
        ]
        with open(file_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(headers)
            for p in products:
                writer.writerow(
                    [
                        p["product_code"],
                        p["medicine_name"],
                        p["generic_name"],
                        p["brand_name"],
                        p["manufacturer"],
                        p["category"],
                        p["dosage_form"],
                        p["strength"],
                        p["pack_size"],
                        p["unit_of_measure"],
                        p["primary_barcode"],
                        str(paisa_to_pkr(p["default_purchase_price_paisa"])),
                        str(paisa_to_pkr(p["default_sale_price_paisa"])),
                        p["min_stock_threshold"],
                        p["sellable_stock"],
                        p["expired_stock"],
                        p["earliest_valid_expiry"] or "",
                        p["verification_status"],
                        p["is_sample_record"],
                        p["is_active"],
                    ]
                )
        return os.path.abspath(file_path)

    def export_inventory_to_excel(self, file_path: str) -> str:
        """Export current medicine master and batch stock to a styled Excel (.xlsx) workbook."""
        products = self.search_products(include_inactive=True)
        os.makedirs(os.path.dirname(os.path.abspath(file_path)), exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Medicine Inventory"
        ws.views.sheetView[0].showGridLines = True

        headers = [
            "Product Code",
            "Medicine Name",
            "Generic / Salt Name",
            "Manufacturer",
            "Category",
            "Dosage Form",
            "Strength",
            "Pack Size",
            "Primary Barcode",
            "Purchase Price (PKR)",
            "Sale Price (PKR)",
            "Min Stock",
            "Sellable Stock",
            "Expired Stock",
            "Earliest Valid Expiry",
            "Verification Status",
        ]
        ws.append(headers)
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=1, column=col_idx)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center", vertical="center")

        for p in products:
            ws.append(
                [
                    p["product_code"],
                    p["medicine_name"],
                    p["generic_name"],
                    p["manufacturer"],
                    p["category"],
                    p["dosage_form"],
                    p["strength"],
                    int(p["pack_size"]),
                    p["primary_barcode"],
                    float(paisa_to_pkr(p["default_purchase_price_paisa"])),
                    float(paisa_to_pkr(p["default_sale_price_paisa"])),
                    int(p["min_stock_threshold"]),
                    int(p["sellable_stock"]),
                    int(p["expired_stock"]),
                    p["earliest_valid_expiry"] or "",
                    p["verification_status"],
                ]
            )
        ws.freeze_panes = "A2"
        wb.save(file_path)
        return os.path.abspath(file_path)

    def import_medicines_from_file(self, user_id: int, file_path: str) -> dict[str, Any]:
        """Import medicines from a validated CSV or Excel (.xlsx) file.

        Returns counts of imported, skipped duplicates, and row-level validation errors.
        """
        ext = os.path.splitext(file_path)[1].lower()
        rows_data: list[dict[str, Any]] = []

        if ext == ".csv":
            with open(file_path, "r", encoding="utf-8-sig") as f:
                reader = csv.DictReader(f)
                for row in reader:
                    rows_data.append({(k or "").strip().lower(): (v or "").strip() for k, v in row.items()})
        elif ext == ".xlsx":
            wb = openpyxl.load_workbook(file_path, data_only=True)
            ws = wb.active
            all_rows = list(ws.iter_rows(values_only=True))
            if all_rows:
                raw_headers = [str(h or "").strip().lower() for h in all_rows[0]]
                for r in all_rows[1:]:
                    if not any(r):
                        continue
                    row_dict = {
                        raw_headers[i]: (str(r[i]).strip() if r[i] is not None else "")
                        for i in range(min(len(raw_headers), len(r)))
                    }
                    rows_data.append(row_dict)
            wb.close()
        else:
            raise ValueError("Unsupported import file format. Please use .csv or .xlsx.")

        imported_count = 0
        skipped_duplicates = 0
        errors: list[str] = []

        for idx, row in enumerate(rows_data, start=2):
            med_name = row.get("medicine_name") or row.get("medicine name") or ""
            if not med_name or med_name.startswith("<--") or "[example]" in med_name.lower():
                continue
            gen_name = row.get("generic_name") or row.get("generic / salt name") or row.get("generic") or ""
            mfg = row.get("manufacturer") or row.get("company") or ""
            category = row.get("category") or "General"
            dosage = row.get("dosage_form") or row.get("dosage form") or "Tablet"
            strength = row.get("strength") or ""
            pack_size_str = row.get("pack_size") or row.get("pack size") or "1"
            unit_str = row.get("unit_of_measure") or row.get("unit") or "Tablet"
            barcode = row.get("barcode") or row.get("primary_barcode") or row.get("primary barcode") or ""
            p_price = row.get("purchase_price_pkr") or row.get("purchase price (pkr)") or "0"
            s_price = row.get("sale_price_pkr") or row.get("sale price (pkr)") or "0"
            min_stock = row.get("min_stock_threshold") or row.get("min stock") or "10"
            batch_no = row.get("batch_number") or row.get("batch") or ""
            expiry = row.get("expiry_date") or row.get("expiry") or ""
            init_qty = row.get("initial_qty") or row.get("qty") or "0"
            ver_status = row.get("verification_status") or "UNVERIFIED_IMPORT"
            if ver_status not in ("VERIFIED", "USER_ENTERED", "SAMPLE_DEMONSTRATION", "UNVERIFIED_IMPORT"):
                ver_status = "UNVERIFIED_IMPORT"
            is_sample = ver_status == "SAMPLE_DEMONSTRATION" or row.get("is_sample_record") in ("1", "true", "True", "YES")

            try:
                self.create_product(
                    user_id=user_id,
                    medicine_name=med_name,
                    generic_name=gen_name,
                    manufacturer=mfg,
                    purchase_price_pkr=p_price,
                    sale_price_pkr=s_price,
                    brand_name=row.get("brand_name") or med_name,
                    category=category,
                    dosage_form=dosage,
                    strength=strength,
                    pack_size=int(float(pack_size_str)),
                    unit_of_measure=unit_str,
                    barcode=barcode if barcode else None,
                    min_stock_threshold=int(float(min_stock)),
                    initial_batch_number=batch_no if batch_no else None,
                    initial_expiry_date=expiry if expiry else None,
                    initial_qty=int(float(init_qty)) if init_qty else 0,
                    is_sample_record=is_sample,
                    verification_status=ver_status,
                    notes=row.get("notes") or "",
                )
                imported_count += 1
            except DuplicateProductError as e:
                skipped_duplicates += 1
                errors.append(f"Row {idx}: Skipped duplicate - {e}")
            except Exception as e:
                errors.append(f"Row {idx} ({med_name}): {e}")

        return {
            "imported_count": imported_count,
            "skipped_duplicates": skipped_duplicates,
            "error_count": len(errors) - skipped_duplicates,
            "messages": errors,
        }

    @staticmethod
    def _validate_date_str(date_str: str) -> None:
        try:
            date.fromisoformat(date_str.strip())
        except ValueError as e:
            raise ValueError(f"Invalid date format '{date_str}'. Expected YYYY-MM-DD.") from e
